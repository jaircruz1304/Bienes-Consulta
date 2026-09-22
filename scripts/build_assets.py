#!/usr/bin/env python3
"""Build data/assets.json from the official Excel workbook.

The workbook remains the source of truth. This script publishes a reduced projection
for fast GitHub Pages queries. It intentionally excludes ID-card and tax-ID fields.
"""
from __future__ import annotations
import argparse, json, os, re, unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook

SHEET_INVENTORY = "Inventario Mejorado"
SHEET_PHOTOS = "Link_fotos"
SHEET_INVOICE_LINKS = "links"
SHEET_DOCUMENTS = "DOCUMENTOS_BIENES"  # optional

ERROR_VALUES = {"#REF!", "#VALUE!", "#N/A", "#NAME?", "VERIFICAR"}

def norm(v):
    s = "" if v is None else str(v)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Z0-9]+", "", s.upper())

def clean(v):
    if v is None: return None
    if isinstance(v, str):
        s=v.strip()
        if not s or s.upper() in ERROR_VALUES: return None
        return s
    return v

def iso(v):
    v=clean(v)
    if isinstance(v, datetime): return v.date().isoformat()
    if isinstance(v, date): return v.isoformat()
    if isinstance(v, str):
        s=v.strip()
        for fmt in ("%d/%m/%Y","%Y-%m-%d","%d-%m-%Y"):
            try: return datetime.strptime(s,fmt).date().isoformat()
            except ValueError: pass
        return s
    return v

def is_url(v):
    if not isinstance(v,str): return False
    try: return urlparse(v.strip()).scheme in {"http","https"}
    except Exception: return False

def rows(ws):
    values=ws.iter_rows(values_only=True)
    header=next(values,None)
    if not header: return []
    keys=[norm(x) for x in header]
    out=[]
    for row in values:
        if not any(clean(v) is not None for v in row): continue
        out.append({keys[i]: clean(row[i]) if i < len(row) else None for i in range(len(keys))})
    return out

def first(r,*names):
    for name in names:
        k=norm(name)
        if clean(r.get(k)) is not None: return clean(r.get(k))
    return None

def starts(r,prefix):
    p=norm(prefix)
    for k,v in r.items():
        if k.startswith(p) and clean(v) is not None: return clean(v)
    return None

def two_col_map(ws):
    m={}
    vals=list(ws.iter_rows(values_only=True))
    for r in vals[1:]:
        if len(r)<2: continue
        k=clean(r[0]); u=clean(r[1])
        if k and is_url(u): m[str(k).strip().upper()]=u.strip()
    return m

def special_statuses(a):
    text=" ".join(str(x or "") for x in [a.get("estado_fisico"),a.get("observaciones"),a.get("motivo_baja")]).upper()
    out=[]
    def add(c,label,level):
        if c not in {x["codigo"] for x in out}: out.append({"codigo":c,"etiqueta":label,"nivel":level})
    if a.get("fecha_baja"): add("BAJA","Dado de baja","critico")
    elif any(x in text for x in ("PARA BAJA","DAR DE BAJA","PROCESO DE BAJA")): add("PROCESO_BAJA","Proceso de baja","advertencia")
    if any(x in text for x in ("MANTENIMIENTO","EN REPARACION","EN REPARACIÓN")): add("MANTENIMIENTO","En mantenimiento","informativo")
    if "SINIESTR" in text: add("SINIESTRO","Siniestrado","critico")
    if any(x in text for x in ("TRANSFERIDO","TRANSFERENCIA REALIZADA")): add("TRANSFERENCIA","Transferido","informativo")
    if any(x in text for x in ("BIEN DONADO","DONACION REALIZADA","DONACIÓN REALIZADA")): add("DONACION","Donado","informativo")
    return out

def doc_rows(wb):
    out={}
    if SHEET_DOCUMENTS not in wb.sheetnames: return out
    for r in rows(wb[SHEET_DOCUMENTS]):
        code=first(r,"CODIGO")
        url=first(r,"URL","URL_DOCUMENTO","LINK")
        if not code or not is_url(url): continue
        vig=first(r,"VIGENTE","ACTIVO")
        if vig is not None and str(vig).strip().upper() in {"NO","0","FALSE","INACTIVO"}: continue
        d={
            "tipo": str(first(r,"TIPO_DOCUMENTO","TIPO") or "OTRO").strip().upper().replace(" ","_"),
            "nombre": first(r,"NOMBRE","DOCUMENTO","DESCRIPCION") or "Documento de respaldo",
            "url": url,
        }
        out.setdefault(str(code).strip().upper(),[]).append(d)
    return out

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--output",default="data/assets.json")
    args=ap.parse_args()

    # data_only=True returns the last values saved by Excel for calculated cells.
    wb=load_workbook(args.workbook,data_only=True,read_only=True)
    if SHEET_INVENTORY not in wb.sheetnames:
        raise SystemExit(f"No existe la hoja requerida: {SHEET_INVENTORY}")

    photos=two_col_map(wb[SHEET_PHOTOS]) if SHEET_PHOTOS in wb.sheetnames else {}
    invoice_links=two_col_map(wb[SHEET_INVOICE_LINKS]) if SHEET_INVOICE_LINKS in wb.sheetnames else {}
    extra_docs=doc_rows(wb)
    assets=[]

    for r in rows(wb[SHEET_INVENTORY]):
        code=first(r,"CODIGO")
        if not code: continue
        code=str(code).strip()
        numero_factura=starts(r,"N FACTURA") or starts(r,"NFACTURA")
        factura_url=first(r,"FACTURA DIGITAL")
        if not is_url(factura_url) and numero_factura:
            factura_url=invoice_links.get(str(numero_factura).strip().upper())
        photo=photos.get(code.upper())
        acta=starts(r,"ACTA ENTREGA")
        poliza=starts(r,"NRO DE POLIZA")

        a={
            "codigo":code,
            "descripcion":first(r,"DESCRIPCION"),
            "descripcion_adicional":starts(r,"DESCRIPCION ADICIONAL"),
            "cantidad":first(r,"CANTIDAD"),
            "tipo_bien":first(r,"TIPO DE BIEN"),
            "clasificacion":starts(r,"ACTIVO FIJO CONTROL ADMINISTRATIVO"),
            "proyecto":starts(r,"PROYECTO FIAS FEIG FAP PASF CONSERVA AVES BIOECONOMIA"),
            "fecha_compra":iso(starts(r,"FECHA DE COMPRA")),
            "donante":first(r,"DONANTE"),
            "proveedor":first(r,"PROVEEDOR"),
            "numero_factura":numero_factura,
            "valor_adquisicion":starts(r,"VALOR DEL BIEN"),
            "marca":first(r,"MARCA"),
            "modelo":first(r,"MODELO"),
            "serie":starts(r,"NUMERO DE SERIE"),
            "custodio":starts(r,"NOMBRES Y APELLIDOS CUSTODIO"),
            "institucion":first(r,"INSTITUCION"),
            "ubicacion":first(r,"UBICACION"),
            "asegurado":starts(r,"ASEGURADO SI NO"),
            "inicio_seguro":iso(first(r,"INICIO SEGURO")),
            "fin_seguro":iso(first(r,"FIN SEGURO")),
            "aseguradora":first(r,"ASEGURADORA"),
            "poliza":poliza,
            "garantia":starts(r,"GARANTIA TECNICA"),
            "inicio_garantia":iso(starts(r,"INICIO GARANTIA")),
            "fin_garantia":iso(starts(r,"FIN DE GARANTIA")),
            "estado_garantia":starts(r,"ESTADO DE GARANTIA"),
            "estado_fisico":starts(r,"ESTADO FISICO DETALLADO"),
            "vida_util":starts(r,"VIDA UTIL ESTIMADA"),
            "fecha_baja":iso(starts(r,"FECHA DE BAJA")),
            "motivo_baja":starts(r,"MOTIVO DE BAJA"),
            "observaciones":first(r,"OBSERVACIONES"),
            "foto_url":photo,
        }
        docs=[]
        if is_url(factura_url): docs.append({"tipo":"FACTURA","nombre":f"Factura {numero_factura or ''}".strip(),"url":factura_url})
        if is_url(acta): docs.append({"tipo":"ACTA_ENTREGA","nombre":"Acta de entrega / custodia","url":acta})
        if is_url(photo): docs.append({"tipo":"FOTOGRAFIA","nombre":"Fotografía del bien","url":photo})
        docs.extend(extra_docs.get(code.upper(),[]))
        # de-duplicate URLs
        seen=set(); a["documentos"]=[]
        for d in docs:
            u=d.get("url")
            if u in seen: continue
            seen.add(u); a["documentos"].append(d)
        a["estados_especiales"]=special_statuses(a)
        assets.append(a)

    codes=[a["codigo"].upper() for a in assets]
    dup=sorted({c for c in codes if codes.count(c)>1})
    if dup: raise SystemExit("Códigos duplicados: "+", ".join(dup[:20]))
    assets.sort(key=lambda x:x["codigo"].upper())

    meta={
        "schemaVersion":2,
        "generatedAt":datetime.now(timezone.utc).isoformat().replace("+00:00","Z"),
        "sourceLastModified":os.environ.get("SOURCE_LAST_MODIFIED") or None,
        "sourceETag":os.environ.get("SOURCE_ETAG") or None,
        "sourceName":os.environ.get("SOURCE_NAME") or Path(args.workbook).name,
        "assetCount":len(assets),
        "mode":"synchronized-json"
    }
    out=Path(args.output);out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps({"meta":meta,"assets":assets},ensure_ascii=False,indent=2),encoding="utf-8")
    Path("data/source-meta.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"OK: {len(assets)} bienes publicados en {out}")

if __name__=="__main__": main()
