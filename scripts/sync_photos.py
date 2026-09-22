#!/usr/bin/env python3
"""Synchronize asset photos from SharePoint/OneDrive into GitHub Pages local assets.

Source of truth:
  Excel sheet ``Link_fotos`` (columns: CÓDIGO, LINK)

Behavior:
- Resolves Microsoft 365 sharing links through Microsoft Graph.
- Accepts mixed original formats (JPEG/JPG, PNG, WEBP, GIF, BMP, TIFF,
  HEIC/HEIF and formats supported by Pillow/pillow-heif).
- Converts every successfully read image to a normalized WEBP file.
- Stores only local public paths in the website; SharePoint links are not needed
  by the visitor's browser.
- Uses an eTag manifest so unchanged photos are not downloaded again.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import os
import re
import sys
import time
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook
from PIL import Image, ImageOps, UnidentifiedImageError

try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except Exception:
    # Standard formats continue to work even if HEIF support is unavailable.
    pass

SHEET_PHOTOS = "Link_fotos"
DEFAULT_MANIFEST = Path("data/photos-manifest.json")
DEFAULT_OUTPUT_DIR = Path("assets/bienes")
GRAPH_BASE = "https://graph.microsoft.com/v1.0"
MAX_EDGE = 2000
WEBP_QUALITY = 86


def clean(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def is_http_url(v: Any) -> bool:
    return isinstance(v, str) and re.match(r"^https?://", v.strip(), re.I) is not None


def safe_code(code: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", code.strip())
    return value.strip(".-") or hashlib.sha256(code.encode("utf-8")).hexdigest()[:16]


def url_hash(url: str) -> str:
    return hashlib.sha256(url.strip().encode("utf-8")).hexdigest()


def version_hash(etag: str | None, blob: bytes | None = None) -> str:
    basis = (etag or "").encode("utf-8")
    if not basis and blob is not None:
        basis = hashlib.sha256(blob).digest()
    return hashlib.sha256(basis).hexdigest()[:14]


def encode_share_url(url: str) -> str:
    token = base64.urlsafe_b64encode(url.encode("utf-8")).decode("ascii").rstrip("=")
    return "u!" + token


class GraphClient:
    def __init__(self, token: str):
        self.session = requests.Session()
        self.session.headers.update({"Authorization": f"Bearer {token}"})

    def request(self, method: str, url: str, **kwargs) -> requests.Response:
        last_exc: Exception | None = None
        for attempt in range(6):
            try:
                r = self.session.request(method, url, timeout=90, **kwargs)
                if r.status_code == 429 or 500 <= r.status_code < 600:
                    retry = r.headers.get("Retry-After")
                    wait = float(retry) if retry and retry.isdigit() else min(2 ** attempt, 30)
                    time.sleep(wait)
                    continue
                r.raise_for_status()
                return r
            except requests.RequestException as exc:
                last_exc = exc
                if attempt >= 5:
                    break
                time.sleep(min(2 ** attempt, 20))
        raise RuntimeError(f"Microsoft Graph no respondió correctamente: {last_exc}")

    def resolve_share(self, share_url: str) -> dict[str, Any]:
        share_id = encode_share_url(share_url)
        endpoint = (
            f"{GRAPH_BASE}/shares/{share_id}/driveItem"
            "?$select=id,name,eTag,parentReference,file,remoteItem"
        )
        meta = self.request("GET", endpoint).json()
        item = meta.get("remoteItem") or meta
        parent = item.get("parentReference") or meta.get("parentReference") or {}
        drive_id = parent.get("driveId")
        item_id = item.get("id") or meta.get("id")
        if not drive_id or not item_id:
            raise RuntimeError("El enlace no resolvió un driveId/itemId de Microsoft 365.")
        if not (item.get("file") or meta.get("file")):
            raise RuntimeError("El enlace no corresponde a un archivo descargable.")
        return {
            "driveId": drive_id,
            "itemId": item_id,
            "name": item.get("name") or meta.get("name") or "imagen",
            "eTag": item.get("eTag") or meta.get("eTag") or "",
        }

    def download(self, drive_id: str, item_id: str) -> bytes:
        endpoint = f"{GRAPH_BASE}/drives/{drive_id}/items/{item_id}/content"
        return self.request("GET", endpoint, allow_redirects=True).content


def normalize_image(blob: bytes, output: Path) -> tuple[int, int]:
    try:
        with Image.open(io.BytesIO(blob)) as im:
            # For animated files we publish a representative first frame.
            try:
                if getattr(im, "is_animated", False):
                    im.seek(0)
            except Exception:
                pass
            im = ImageOps.exif_transpose(im)
            has_alpha = "A" in im.getbands() or "transparency" in im.info
            im = im.convert("RGBA" if has_alpha else "RGB")
            im.thumbnail((MAX_EDGE, MAX_EDGE), Image.Resampling.LANCZOS)
            width, height = im.size
            output.parent.mkdir(parents=True, exist_ok=True)
            tmp = output.with_suffix(output.suffix + ".tmp")
            im.save(tmp, format="WEBP", quality=WEBP_QUALITY, method=6)
            os.replace(tmp, output)
            return width, height
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise RuntimeError(f"Formato de imagen no reconocido o dañado: {exc}") from exc


def load_manifest(path: Path) -> dict[str, Any]:
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(obj, dict) and isinstance(obj.get("assets"), dict):
            return obj
    except Exception:
        pass
    return {"schemaVersion": 1, "assets": {}}


def read_photo_sources(workbook: str) -> dict[str, str]:
    wb = load_workbook(workbook, data_only=True, read_only=True)
    if SHEET_PHOTOS not in wb.sheetnames:
        return {}
    ws = wb[SHEET_PHOTOS]
    result: dict[str, str] = {}
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:
            continue
        if len(row) < 2:
            continue
        code, url = clean(row[0]), clean(row[1])
        if code and is_http_url(url):
            result[code.upper()] = url.strip()
    return result


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--manifest", default=str(DEFAULT_MANIFEST))
    ap.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    ap.add_argument("--refresh-all", action="store_true", help="Check source eTags even when the Excel link did not change.")
    args = ap.parse_args()

    token = os.environ.get("GRAPH_TOKEN", "").strip()
    if not token:
        raise SystemExit("GRAPH_TOKEN no está disponible.")

    manifest_path = Path(args.manifest)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    old = load_manifest(manifest_path)
    old_assets: dict[str, Any] = old.get("assets", {})
    sources = read_photo_sources(args.workbook)
    graph = GraphClient(token)

    new_assets: dict[str, Any] = {}
    downloaded = 0
    reused = 0
    warnings = 0
    removed = 0

    for code, share_url in sorted(sources.items()):
        key = code.upper()
        previous = old_assets.get(key) if isinstance(old_assets.get(key), dict) else None
        h = url_hash(share_url)
        filename = safe_code(code) + ".webp"
        local_file = output_dir / filename
        local_path = "./" + local_file.as_posix()

        same_source = bool(previous and previous.get("urlHash") == h)
        can_reuse_without_graph = (
            not args.refresh_all
            and same_source
            and previous.get("status") == "ok"
            and local_file.exists()
        )
        if can_reuse_without_graph:
            new_assets[key] = previous
            reused += 1
            continue

        try:
            resolved = graph.resolve_share(share_url)
            etag = str(resolved.get("eTag") or "")
            if (
                same_source
                and previous
                and previous.get("status") == "ok"
                and previous.get("sourceETag") == etag
                and local_file.exists()
            ):
                new_assets[key] = previous
                reused += 1
                continue

            blob = graph.download(resolved["driveId"], resolved["itemId"])
            if not blob:
                raise RuntimeError("El archivo descargado está vacío.")
            width, height = normalize_image(blob, local_file)
            new_assets[key] = {
                "status": "ok",
                "localPath": local_path,
                "version": version_hash(etag, blob),
                "sourceETag": etag,
                "sourceName": resolved.get("name") or "imagen",
                "urlHash": h,
                "width": width,
                "height": height,
                "format": "webp",
            }
            downloaded += 1
            print(f"[OK] {code}: {resolved.get('name')} -> {local_path}")
        except Exception as exc:
            warnings += 1
            # If the source link itself did not change, keep the last valid local copy
            # during a temporary Graph/SharePoint failure.
            if same_source and previous and previous.get("status") == "ok" and local_file.exists():
                new_assets[key] = previous
                reused += 1
                print(f"[WARN] {code}: no se pudo refrescar; se conserva la copia local anterior. {exc}", file=sys.stderr)
            else:
                if local_file.exists():
                    local_file.unlink()
                new_assets[key] = {
                    "status": "error",
                    "urlHash": h,
                    "message": str(exc)[:300],
                }
                print(f"[WARN] {code}: {exc}", file=sys.stderr)

    # Remove files for codes that no longer have a photo link in the official Excel.
    active = set(sources)
    for code, info in old_assets.items():
        if code in active or not isinstance(info, dict):
            continue
        p = info.get("localPath")
        if isinstance(p, str) and p.startswith("./assets/bienes/"):
            fp = Path(p[2:])
            if fp.exists() and fp.is_file():
                fp.unlink()
                removed += 1

    manifest = {"schemaVersion": 2, "assets": new_assets}
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    new_text = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    old_text = manifest_path.read_text(encoding="utf-8") if manifest_path.exists() else ""
    if new_text != old_text:
        manifest_path.write_text(new_text, encoding="utf-8")

    ok_count = sum(1 for v in new_assets.values() if isinstance(v, dict) and v.get("status") == "ok")
    print(
        f"Fotografías: fuentes={len(sources)}, locales_ok={ok_count}, "
        f"descargadas/actualizadas={downloaded}, reutilizadas={reused}, eliminadas={removed}, avisos={warnings}"
    )

    if sources and ok_count == 0:
        print("ERROR: no fue posible generar ninguna fotografía local.", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
